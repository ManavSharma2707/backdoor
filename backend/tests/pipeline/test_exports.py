"""PDF, Excel and CSV exports: structure, the Rupee sign, formulas and totals to the paisa."""

import copy
import csv
import io
import re
import uuid
from decimal import Decimal
from typing import Any

import openpyxl
import pytest
from pypdf import PdfReader
from sqlalchemy import text

from app.services.export.common import ExportData, ExportLine, inr, load_export
from app.services.export.csv_export import PLAN_COLUMNS, PURCHASE_COLUMNS, build_csv
from app.services.export.pdf_export import build_pdf
from app.services.export.xlsx_export import SHEETS, build_xlsx
from tests.pipeline.conftest import Api

PAISA = Decimal("0.01")


def download(api_: Api, fmt: str, **params: Any) -> Any:
    return api_.get(f"/reports/{api_.report_a}/export", params={"format": fmt, **params})


def workbook(content: bytes, *, values: bool) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content), data_only=values)


def money(value: Any) -> Decimal:
    return Decimal(str(value)).quantize(PAISA)


def csv_rows(content: bytes) -> list[dict[str, str]]:
    assert content.startswith(b"\xef\xbb\xbf"), "the file needs a byte order mark"
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))


def approve_some(api_: Api, n: int = 5) -> list[dict[str, Any]]:
    funded = [
        ln
        for ln in api_.lines(section="buy")
        if Decimal(ln["allocated_qty"]) > 0 and ln["unit_cost"] and ln["supplier_id"]
    ][:n]
    for i, ln in enumerate(funded):
        if i == 1:
            pack = Decimal(ln["pack_size"])
            body = {"decision": "edited", "final_qty": str(pack), "reason_code": "shelf_space"}
        else:
            body = {"decision": "approved"}
        assert api_.put(f"/report-lines/{ln['id']}/decision", body).status_code == 200
    return funded


# --------------------------------------------------------------------------- the endpoint


def test_headers_names_and_authorisation(api: Api) -> None:
    for fmt, extension, media in (
        ("pdf", "pdf", "application/pdf"),
        ("xlsx", "xlsx", "spreadsheetml.sheet"),
        ("csv", "csv", "text/csv"),
    ):
        response = download(api, fmt)
        assert response.status_code == 200
        assert media in response.headers["content-type"]
        disposition = response.headers["content-disposition"]
        assert disposition == (
            f'attachment; filename="backroom_shree-sai-kirana_monday-plan_2026-10-12.{extension}"'
        )
        assert response.headers["cache-control"] == "no-store"
    lines_file = download(api, "csv", kind="plan_lines")
    assert "monday-plan-lines_2026-10-12.csv" in lines_file.headers["content-disposition"]
    assert (
        api.get(
            f"/reports/{api.report_a}/export", headers=api.b, params={"format": "pdf"}
        ).status_code
        == 404
    )
    assert download(api, "docx").status_code == 422
    with api.db.session() as s:
        n = s.scalar(text("SELECT count(*) FROM audit_logs WHERE action = 'report.export'"))
        assert (n or 0) >= 4


# --------------------------------------------------------------------------- csv


def test_purchase_list_csv(api: Api) -> None:
    everything = csv_rows(download(api, "csv").content)
    assert list(everything[0]) == list(PURCHASE_COLUMNS)
    funded = [
        ln
        for ln in api.lines(section="buy")
        if Decimal(ln["allocated_qty"]) > 0 and ln["unit_cost"] and ln["supplier_id"]
    ]
    assert len(everything) == len(funded)
    assert {r["decision"] for r in everything} == {"pending"}
    by_sku = {r["sku"]: r for r in everything}
    for ln in funded:
        row = by_sku[ln["sku"]]
        assert Decimal(row["quantity"]) == Decimal(ln["allocated_qty"])
        assert "," not in row["unit_cost"]
        assert money(row["line_total"]) == money(ln["funded_spend"])
        assert re.fullmatch(r"\d{4}-\d\d-\d\d", row["expected_delivery_date"])
    # With decisions only approved and edited lines appear, at the decided quantity.
    chosen = approve_some(api, 5)
    decided = csv_rows(download(api, "csv").content)
    assert {r["sku"] for r in decided} == {ln["sku"] for ln in chosen}
    edited = next(r for r in decided if r["decision"] == "edited")
    source = next(ln for ln in chosen if ln["sku"] == edited["sku"])
    assert Decimal(edited["quantity"]) == Decimal(source["pack_size"])
    assert Decimal(edited["packs"]) == 1
    assert money(edited["line_total"]) == money(
        Decimal(source["pack_size"]) * Decimal(source["unit_cost"])
    )
    assert edited["reason"] == "shelf_space"
    only = chosen[0]["supplier_id"]
    narrowed = csv_rows(download(api, "csv", supplier_id=only).content)
    assert {r["supplier_code"] for r in narrowed} == {
        next(r["supplier_code"] for r in decided if r["sku"] == chosen[0]["sku"])
    }


def test_plan_lines_csv_has_every_line_and_machine_names(api: Api) -> None:
    rows = csv_rows(download(api, "csv", kind="plan_lines").content)
    assert list(rows[0]) == list(PLAN_COLUMNS)
    assert len(rows) == 92
    first = rows[0]
    assert first["sku"] and first["product_name"] and first["reason_text"]
    assert all(r["risk_level"] in ("critical", "high", "medium", "low", "none") for r in rows)
    ids = {ln["sku"] for ln in api.lines()}
    assert {r["sku"] for r in rows} == ids


def test_csv_cells_cannot_become_formulas(api: Api) -> None:
    data = load_export(api.db.session(), _ctx(api), uuid.UUID(api.report_a))
    data.lines[0].data["name"] = '=HYPERLINK("http://evil.example","click")'
    data.lines[1].data["name"] = "+1+1"
    text_out = build_csv(data, "plan_lines").decode("utf-8-sig")
    assert "'=HYPERLINK" in text_out
    assert "'+1+1" in text_out


def _ctx(api_: Api) -> Any:
    from app.core.audit import RequestMeta
    from app.core.deps import TenantContext

    with api_.db.session() as s:
        user = s.execute(
            text("SELECT id FROM users WHERE tenant_id = :t"), {"t": api_.db.tenant_id("fmcg_a")}
        ).scalar_one()
    return TenantContext(api_.db.tenant_id("fmcg_a"), user, "owner", RequestMeta())


# --------------------------------------------------------------------------- excel


def test_workbook_structure_formats_and_formulas(api: Api) -> None:
    content = download(api, "xlsx").content
    wb = workbook(content, values=False)
    assert wb.sheetnames == list(SHEETS)
    buy = wb["Buy list"]
    assert buy.freeze_panes == "A2"
    assert buy.auto_filter.ref and buy.auto_filter.ref.startswith("A1:")
    header = [c.value for c in buy[1]]
    assert header[:3] == ["Supplier", "SKU", "Product"]
    n_buy = len(api.lines(section="buy"))
    first_data = [c.value for c in buy[2]]
    assert first_data[15] == "=N2*O2"  # funded spend is quantity times cost
    total_row = n_buy + 2
    assert buy.cell(row=total_row, column=16).value == f"=SUBTOTAL(109,P2:P{n_buy + 1})"
    assert buy.cell(row=2, column=15).number_format == "[$₹-4009]#,##0.00"
    assert buy.cell(row=2, column=5).number_format == "#,##0.###"
    assert buy.cell(row=2, column=15).data_type == "n"
    for name in SHEETS:
        ws = wb[name]
        if name not in ("Summary", "By supplier", "Budget", "Accuracy"):
            assert ws.freeze_panes == "A2", name
    values = workbook(content, values=True)
    vbuy = values["Buy list"]
    summary = api.get(f"/reports/{api.report_a}").json()["summary"]
    assert money(vbuy.cell(row=total_row, column=16).value) == money(summary["planned_spend"])
    assert vbuy.cell(row=total_row + 3, column=16).value == "Matches the plan"
    sheet_summary = values["Summary"]
    labels = {r[0].value: r[1].value for r in sheet_summary.iter_rows(min_row=1, max_row=40)}
    assert labels["Planned spend (funded)"] == pytest.approx(float(summary["planned_spend"]))
    assert "Demo data. Numbers are simulated." in labels
    assert labels["Model"] == "forecast-engine 1.0.0"


def test_every_formula_has_a_cached_value(api: Api) -> None:
    content = download(api, "xlsx").content
    formulas = workbook(content, values=False)
    cached = workbook(content, values=True)
    checked = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cached[ws.title][cell.coordinate].value is not None, (
                        ws.title,
                        cell.coordinate,
                    )
                    checked += 1
    assert checked > 30


def test_text_starting_with_equals_is_stored_as_text(api: Api) -> None:
    data = load_export(api.db.session(), _ctx(api), uuid.UUID(api.report_a))
    data.lines[0].data["name"] = '=HYPERLINK("http://evil.example","click")'
    bad = next(ln for ln in data.section_lines("buy"))
    bad.data["name"] = "=1+1"
    wb = workbook(build_xlsx(data), values=False)
    cell = wb["Buy list"]["C2"]
    assert cell.data_type == "s"
    assert cell.value == "=1+1"


def test_by_supplier_sheet_groups_with_subtotals(api: Api) -> None:
    approve_some(api, 6)
    wb = workbook(download(api, "xlsx").content, values=True)
    ws = wb["By supplier"]
    subtotals = [row for row in ws.iter_rows(values_only=True) if row[5] == "Subtotal"]
    assert subtotals
    total = sum(money(r[6]) for r in subtotals)
    purchase = csv_rows(download(api, "csv").content)
    assert total == sum(money(r["line_total"]) for r in purchase)


# --------------------------------------------------------------------------- pdf


def pdf_text(content: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() for page in reader.pages), len(reader.pages)


def test_pdf_has_the_business_week_demo_label_and_rupee_sign(api: Api) -> None:
    content = download(api, "pdf").content
    assert content.startswith(b"%PDF")
    text_out, pages = pdf_text(content)
    assert 3 <= pages <= 12
    assert "Shree Sai Kirana and General Stores" in text_out
    assert "Monday plan, week of 12 Oct 2026" in text_out
    assert "Demo data. Numbers are simulated." in text_out
    assert "₹" in text_out
    assert "Data up to 11 Oct 2026" in text_out
    assert "Page 1 of" in text_out
    for heading in (
        "Buy this week",
        "Do not buy yet",
        "Watch closely",
        "Reduce exposure",
        "Festival and season",
        "Budget view",
        "Data and confidence notes",
        "How to read this plan",
    ):
        assert heading in text_out
    summary = api.get(f"/reports/{api.report_a}").json()["summary"]
    assert inr(summary["planned_spend"]) in text_out
    assert not re.search(r"[–—]", text_out)


def test_pdf_section_and_supplier_filters(api: Api) -> None:
    _, full_pages = pdf_text(download(api, "pdf").content)
    only, only_pages = pdf_text(download(api, "pdf", section="watch").content)
    assert "Watch closely" in only
    assert "Buy this week" not in only
    assert only_pages < full_pages


def test_pdf_with_150_lines_still_renders(api: Api) -> None:
    data = load_export(api.db.session(), _ctx(api), uuid.UUID(api.report_a))
    buy_ids = data.payload["sections"]["buy"]
    template = data.by_id[buy_ids[0]]
    for i in range(150):
        clone = copy.deepcopy(template.data)
        clone["id"] = str(uuid.uuid4())
        clone["sku"] = f"X-{i:03d}"
        clone["name"] = f"Test product number {i}"
        data.by_id[clone["id"]] = ExportLine(clone, "pending", None, None)
        data.payload["sections"]["buy"].append(clone["id"])
    text_out, pages = pdf_text(build_pdf(data))
    assert pages <= 20
    assert "Test product number 149" in text_out


def test_money_formatting_uses_indian_grouping() -> None:
    assert inr("182340") == "₹1,82,340"
    assert inr("1234567.5", 2) == "₹12,34,567.50"
    assert inr("999") == "₹999"
    assert inr("-45000") == "₹-45,000"
    assert inr(None) == ""


# --------------------------------------------------------------------------- across formats


def test_totals_agree_across_every_format_to_the_paisa(api: Api) -> None:
    summary = api.get(f"/reports/{api.report_a}").json()["summary"]
    planned = money(summary["planned_spend"])
    # Without decisions: the funded lines.
    purchase = csv_rows(download(api, "csv").content)
    csv_total = sum(money(r["line_total"]) for r in purchase)
    wb = workbook(download(api, "xlsx").content, values=True)
    n_buy = len(api.lines(section="buy"))
    xlsx_total = money(wb["Buy list"].cell(row=n_buy + 2, column=16).value)
    by_supplier = sum(
        money(r[6]) for r in wb["By supplier"].iter_rows(values_only=True) if r[5] == "Subtotal"
    )
    plan_rows = csv_rows(download(api, "csv", kind="plan_lines").content)
    plan_total = sum(money(r["funded_spend"]) for r in plan_rows if r["funded_spend"])
    pdf_text_out, _ = pdf_text(download(api, "pdf").content)
    assert csv_total == xlsx_total == by_supplier == planned
    assert plan_total == planned
    assert f"Total funded: {inr(planned)}" in pdf_text_out
    # With decisions the purchase list follows the decided quantities in every format.
    approve_some(api, 6)
    decided_csv = csv_rows(download(api, "csv").content)
    decided_total = sum(money(r["line_total"]) for r in decided_csv)
    wb2 = workbook(download(api, "xlsx").content, values=True)
    subtotal_total = sum(
        money(r[6]) for r in wb2["By supplier"].iter_rows(values_only=True) if r[5] == "Subtotal"
    )
    assert decided_total == subtotal_total
    assert decided_total < planned


def test_every_exported_line_exists_in_the_plan(api: Api) -> None:
    lines = {ln["sku"] for ln in api.lines()}
    for row in csv_rows(download(api, "csv", kind="plan_lines").content):
        assert row["sku"] in lines
    wb = workbook(download(api, "xlsx").content, values=True)
    for row in wb["Buy list"].iter_rows(
        min_row=2, max_row=len(api.lines(section="buy")) + 1, values_only=True
    ):
        assert row[1] in lines
    assert isinstance(ExportData, type)
